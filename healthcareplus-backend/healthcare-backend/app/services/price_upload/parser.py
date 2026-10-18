"""Parse + validate price-upload files. No database access here.

Two layouts, one output (`ParsedFile.records`):
  * wide : col A = test name, other headers = "LabName(PinCode)", cells = price
  * long : columns lab_name, pin_code, test_name, price
"""
from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable, Iterator

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HEADER_RE = re.compile(r"^(.+?)\s*\((\d{6})\)$", re.ASCII)
MAX_PRICE = 10_000_000.0  # catches typos like an extra "0000"
LONG_COLUMNS = ("lab_name", "pin_code", "test_name", "price")


# ---------- small helpers -------------------------------------------------
def normalize_name(value: Any) -> str:
    """trim, lowercase, collapse inner whitespace."""
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def clean_text(value: Any) -> str:
    """trim + collapse whitespace, keep original case (for display)."""
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def parse_header(header: Any) -> tuple[str, str] | None:
    """'Apollo (800012)' -> ('Apollo', '800012'); None if malformed."""
    m = HEADER_RE.match(clean_text(header))
    if not m:
        return None
    name = m.group(1).strip()
    return (name, m.group(2)) if name else None


def header_error_reason(header: Any) -> str:
    text = clean_text(header)
    if not text:
        return "Header is empty but the column has data"
    if "(" not in text or not text.endswith(")"):
        return f'Header "{text}" must look like LabName(PinCode), e.g. Apollo(800012) — missing bracket'
    inside = text[text.rfind("(") + 1 : -1].strip()
    if not inside.isdigit():
        return f'Header "{text}" has a non-numeric pin code "{inside}" — pin must be 6 digits'
    if len(inside) != 6:
        return f'Header "{text}" has a {len(inside)}-digit pin code — pin must be exactly 6 digits'
    return f'Header "{text}" has no lab name before the pin code'


def parse_price(value: Any) -> tuple[float | None, str | None]:
    """Return (price, None) or (None, reason). Blank handled by the caller."""
    if isinstance(value, bool):
        return None, f"Price {value!r} is not a number"
    if isinstance(value, (int, float, Decimal)):
        raw: Any = value
    else:
        text = str(value).strip().replace(",", "")
        raw = text
    try:
        d = Decimal(str(raw))
    except (InvalidOperation, ValueError):
        return None, f'Price "{value}" is not a number'
    if not d.is_finite():
        return None, f'Price "{value}" is not a number'
    price = float(d)
    if math.isnan(price) or math.isinf(price):
        return None, f'Price "{value}" is not a number'
    if price <= 0:
        return None, f"Price {value} must be greater than 0"
    if price > MAX_PRICE:
        return None, f"Price {value} is unrealistically large (max {int(MAX_PRICE)})"
    return round(price, 2), None


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


# ---------- result types --------------------------------------------------
@dataclass
class RowError:
    row: int
    column: str
    reason: str


@dataclass
class PriceRecord:
    row: int
    column: str
    lab_name: str
    pin_code: str
    test_name: str
    price: float

    @property
    def lab_key(self) -> tuple[str, str]:
        return (normalize_name(self.lab_name), self.pin_code)

    @property
    def test_key(self) -> str:
        return normalize_name(self.test_name)


@dataclass
class ParsedFile:
    layout: str  # "wide" | "long"
    records: list[PriceRecord] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    branch_count: int = 0
    row_count: int = 0


class FileFormatError(ValueError):
    """File is unreadable / wrong type — nothing could be parsed."""


# ---------- reading -------------------------------------------------------
def _read_csv(raw: bytes) -> Iterator[list[Any]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    yield from csv.reader(io.StringIO(text))


def _read_xlsx(raw: bytes) -> Iterator[list[Any]]:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # zip/xml errors
        raise FileFormatError("Could not read the Excel file — is it a valid .xlsx?") from exc
    ws = wb.worksheets[0]
    for row in ws.iter_rows(values_only=True):
        yield list(row)
    wb.close()


def read_rows(filename: str, raw: bytes) -> list[list[Any]]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        rows = list(_read_csv(raw))
    elif name.endswith(".xlsx"):
        rows = list(_read_xlsx(raw))
    else:
        raise FileFormatError("Only .csv and .xlsx files are accepted.")
    # drop trailing fully-blank rows but keep row numbering intact
    while rows and all(_is_blank(c) for c in rows[-1]):
        rows.pop()
    if not rows:
        raise FileFormatError("The file is empty.")
    return rows


# ---------- entry point ---------------------------------------------------
def parse_file(filename: str, raw: bytes) -> ParsedFile:
    rows = read_rows(filename, raw)
    header = [clean_text(c).lower() for c in rows[0]]
    if all(col in header for col in LONG_COLUMNS):
        return parse_long(rows)
    return parse_wide(rows)


# ---------- wide ----------------------------------------------------------
def parse_wide(rows: list[list[Any]]) -> ParsedFile:
    result = ParsedFile(layout="wide")
    header = rows[0]
    body = rows[1:]
    result.row_count = len(body)

    # Resolve each lab column once.
    col_labs: dict[int, tuple[str, str]] = {}  # col idx -> (display name, pin)
    seen_keys: dict[tuple[str, str], str] = {}  # lab_key -> first column letter
    for idx in range(1, len(header)):
        letter = get_column_letter(idx + 1)
        h = header[idx] if idx < len(header) else None
        has_data = any(idx < len(r) and not _is_blank(r[idx]) for r in body)
        if _is_blank(h):
            if has_data:
                result.errors.append(RowError(1, letter, header_error_reason(h)))
            continue
        parsed = parse_header(h)
        if parsed is None:
            result.errors.append(RowError(1, letter, header_error_reason(h)))
            continue
        key = (normalize_name(parsed[0]), parsed[1])
        if key in seen_keys:
            result.errors.append(
                RowError(
                    1,
                    letter,
                    f'Duplicate branch "{parsed[0]}({parsed[1]})" — already in column {seen_keys[key]}',
                )
            )
            continue
        seen_keys[key] = letter
        col_labs[idx] = parsed
    result.branch_count = len(col_labs)

    seen_tests: dict[str, int] = {}
    for offset, row in enumerate(body):
        rownum = offset + 2
        if all(_is_blank(c) for c in row):
            continue
        test_raw = row[0] if row else None
        test_name = clean_text(test_raw)
        priced = [i for i in col_labs if i < len(row) and not _is_blank(row[i])]
        if not test_name:
            if priced:
                result.errors.append(RowError(rownum, "A", "Test name is empty but the row has prices"))
            continue
        tkey = normalize_name(test_name)
        if tkey in seen_tests:
            if priced:
                result.errors.append(
                    RowError(rownum, "A", f'Duplicate test row "{test_name}" — first seen on row {seen_tests[tkey]}')
                )
            continue
        seen_tests[tkey] = rownum
        for idx in priced:
            letter = get_column_letter(idx + 1)
            price, reason = parse_price(row[idx])
            if reason:
                result.errors.append(RowError(rownum, letter, reason))
                continue
            lab_name, pin = col_labs[idx]
            result.records.append(
                PriceRecord(rownum, letter, lab_name, pin, test_name, price)  # type: ignore[arg-type]
            )
    return result


# ---------- long ----------------------------------------------------------
def parse_long(rows: list[list[Any]]) -> ParsedFile:
    result = ParsedFile(layout="long")
    header = [clean_text(c).lower() for c in rows[0]]
    idx = {c: header.index(c) for c in LONG_COLUMNS}
    letters = {c: get_column_letter(idx[c] + 1) for c in LONG_COLUMNS}
    body = rows[1:]
    result.row_count = len(body)

    seen: dict[tuple[tuple[str, str], str], int] = {}
    branches: set[tuple[str, str]] = set()

    def cell(row: list[Any], col: str) -> Any:
        i = idx[col]
        return row[i] if i < len(row) else None

    for offset, row in enumerate(body):
        rownum = offset + 2
        if all(_is_blank(c) for c in row):
            continue
        lab = clean_text(cell(row, "lab_name"))
        pin_raw = cell(row, "pin_code")
        pin = clean_text(pin_raw)
        if isinstance(pin_raw, float) and pin_raw.is_integer():
            pin = str(int(pin_raw))  # Excel gave 800012.0 — never accept as-is otherwise
        test = clean_text(cell(row, "test_name"))
        price_raw = cell(row, "price")

        problems = False
        if not lab:
            result.errors.append(RowError(rownum, letters["lab_name"], "Lab name is empty"))
            problems = True
        if not re.fullmatch(r"[0-9]{6}", pin):
            result.errors.append(
                RowError(rownum, letters["pin_code"], f'Pin code "{pin}" must be exactly 6 digits')
            )
            problems = True
        if not test:
            result.errors.append(RowError(rownum, letters["test_name"], "Test name is empty"))
            problems = True
        if _is_blank(price_raw):
            # long format has no "empty = unavailable" meaning; a row without price is a mistake
            result.errors.append(RowError(rownum, letters["price"], "Price is empty"))
            problems = True
            price = None
        else:
            price, reason = parse_price(price_raw)
            if reason:
                result.errors.append(RowError(rownum, letters["price"], reason))
                problems = True
        if problems:
            continue

        key = ((normalize_name(lab), pin), normalize_name(test))
        if key in seen:
            result.errors.append(
                RowError(rownum, letters["price"], f"Duplicate of row {seen[key]} (same branch + test)")
            )
            continue
        seen[key] = rownum
        branches.add(key[0])
        result.records.append(PriceRecord(rownum, letters["price"], lab, pin, test, price))  # type: ignore[arg-type]
    result.branch_count = len(branches)
    return result


# ---------- master test list (one column of names) ------------------------
@dataclass
class TestListParsed:
    names: list[str] = field(default_factory=list)  # unique by normalized name, first spelling kept
    total_rows: int = 0
    blank_rows: int = 0
    duplicates_in_file: int = 0
    too_long: list[int] = field(default_factory=list)  # row numbers


def parse_test_list(filename: str, raw: bytes, max_len: int = 500) -> TestListParsed:
    """First column = test names. A header like 'Test Name' is skipped."""
    rows = read_rows(filename, raw)
    out = TestListParsed()
    first = clean_text(rows[0][0] if rows[0] else "").lower()
    start = 1 if first in {"test name", "test", "name", "test names", "tests"} else 0
    seen: set[str] = set()
    for offset, row in enumerate(rows[start:]):
        out.total_rows += 1
        name = clean_text(row[0] if row else "")
        if not name:
            out.blank_rows += 1
            continue
        if len(name) > max_len:
            out.too_long.append(offset + start + 1)
            continue
        key = normalize_name(name)
        if key in seen:
            out.duplicates_in_file += 1
            continue
        seen.add(key)
        out.names.append(name)
    return out
